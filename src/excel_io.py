"""Импорт активного листа Excel и экспорт полной спецификации."""

from decimal import Decimal

from calculations import MAX_ROWS
from utils import decimal_to_float, match_field, parse_decimal, quantize_money

FIELDS = ("name", "gost", "unit", "qty", "net", "gross", "vat", "sum_net", "sum_gross")
HEADERS = [
    "№",
    "Наименование товара",
    "ГОСТ / ТУ",
    "Ед. изм.",
    "Количество",
    "Цена без НДС",
    "Цена с НДС",
    "НДС за единицу",
    "Сумма без НДС",
    "Сумма с НДС",
]


def read_excel(path):
    """Находит шапку в первых 15 строках; пропускает пустые строки и итоги.

    Формулы импортируются по последнему значению, сохранённому Excel.
    Без шапки ожидаются колонки: наименование, цена без НДС, цена с НДС, НДС.
    """
    from openpyxl import load_workbook

    workbook = load_workbook(path, data_only=True, read_only=True)
    try:
        sheet = workbook.active
        if sheet is None:
            raise ValueError("В книге нет активного листа.")
        if (sheet.max_row or 0) > 10000:
            raise ValueError("Лист слишком большой: допустимо до 10 000 строк, включая пустые.")
        preview = list(
            sheet.iter_rows(
                min_row=1,
                max_row=min(sheet.max_row or 1, 15),
                max_col=min(sheet.max_column or 1, 50),
                values_only=True,
            )
        )
        best_row, mapping = 0, {}
        for index, values in enumerate(preview, 1):
            candidate = {}
            for column, value in enumerate(values):
                field = match_field(value)
                if field and field not in candidate:
                    candidate[field] = column
            if "name" in candidate and len(candidate) > len(mapping):
                best_row, mapping = index, candidate
        if len(mapping) < 2:
            best_row, mapping = 0, {"name": 0, "net": 1, "gross": 2, "vat": 3}
        rows = []
        for values in sheet.iter_rows(
            min_row=best_row + 1, max_col=max(mapping.values()) + 1, values_only=True
        ):

            def get(key):
                column = mapping.get(key)
                return values[column] if column is not None else None

            # Наш экспорт ставит «Итого» в первую колонку, сторонние — в наименование.
            first = str(values[0] or "").strip().lower().rstrip(":")
            name = str(get("name") or "").strip()
            if first in {"итого", "всего", "total"} or name.lower().rstrip(":") in {
                "итого",
                "всего",
                "total",
            }:
                continue
            row = {key: "" if get(key) is None else str(get(key)).strip() for key in FIELDS}
            if not any(row.values()):
                continue
            # Не маскируем испорченную цену пустым значением: проверит calculations.
            for key in ("qty", "net", "gross", "vat", "sum_net", "sum_gross"):
                value = parse_decimal(row[key])
                if value is not None:
                    row[key] = str(value)
            rows.append(row)
            if len(rows) > MAX_ROWS:
                raise ValueError(f"Допускается не более {MAX_ROWS} товарных строк.")
        return rows
    finally:
        workbook.close()


def _sum_rows(rows, key):
    return quantize_money(sum((row.get(key) or Decimal(0) for row in rows), Decimal(0)))


def write_excel(path, vat_rate, rows):
    """Записывает цены, количество и суммы; названия хранятся как текст, не формулы."""
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Спецификация"
    sheet.append(HEADERS)
    sheet["H1"].comment = Comment(f"Ставка НДС: {vat_rate}%. НДС за единицу товара.", "PyNDS")
    for index, row in enumerate(rows, 1):
        values = [index] + [
            row[key] if key in {"name", "gost", "unit"} else decimal_to_float(row.get(key))
            for key in FIELDS
        ]
        sheet.append(values)
        for column in (2, 3, 4):
            sheet.cell(sheet.max_row, column).data_type = "s"
    sheet.append(
        [
            "Итого",
            "",
            "",
            "",
            "",
            "",
            "",
            "",
            decimal_to_float(_sum_rows(rows, "sum_net")),
            decimal_to_float(_sum_rows(rows, "sum_gross")),
        ]
    )
    border = Border(*([Side(style="thin", color="E5D9CC")] * 4))
    for row in sheet:
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            if cell.row in (1, sheet.max_row):
                cell.font = Font(bold=True, color="243E35")
                cell.fill = PatternFill("solid", fgColor="F4ECE2")
            if cell.row > 1 and cell.column >= 6:
                cell.number_format = "#,##0.00"
                cell.alignment = Alignment(horizontal="right", vertical="center")
            elif cell.row > 1 and cell.column == 5:
                cell.number_format = "0.###"
    for column, width in zip("ABCDEFGHIJ", (7, 38, 18, 12, 14, 19, 19, 19, 20, 20)):
        sheet.column_dimensions[column].width = width
    sheet.row_dimensions[1].height = 34
    sheet.freeze_panes = "C2"
    sheet.auto_filter.ref = f"A1:J{len(rows) + 1}"
    sheet.print_title_rows = "1:1"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    workbook.save(path)
    workbook.close()
