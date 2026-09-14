"""Экспорт спецификации в Word."""

from decimal import Decimal

from settings_store import load_spec_context
from utils import format_decimal, format_quantity, quantize_money


def write_word(path, vat_rate, rows, settings=None):
    """Создает Word-документ спецификации."""
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Mm, Pt

    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Mm(297)
    section.page_height = Mm(210)
    # Поля как в шаблоне "Узкие".
    section.left_margin = Cm(1.27)
    section.right_margin = Cm(1.27)
    section.top_margin = Cm(1.27)
    section.bottom_margin = Cm(1.27)

    base_style = doc.styles["Normal"]
    base_style.font.name = "Times New Roman"
    base_style.font.size = Pt(12)

    title = doc.add_paragraph("Спецификация")
    title_run = title.runs[0]
    title_run.bold = True
    title_run.font.size = Pt(14)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    spec_context = load_spec_context(settings)

    # Таблица с фиксированными колонками под стиль спецификации.
    table = doc.add_table(rows=0, cols=12)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = False

    # Задаём сетку до объединения: ширина объединённой ячейки равна сумме колонок.
    weights = [0.8, 3.0, 3.0, 1.2, 1.2, 1.2, 1.4, 2.5, 1.9, 2.5, 2.8, 2.8]
    available = section.page_width - section.left_margin - section.right_margin
    for column, weight in zip(table.columns, weights):
        column.width = int(available * weight / sum(weights))

    def set_cell(
        cell,
        text,
        bold=False,
        align=WD_ALIGN_PARAGRAPH.LEFT,
        valign=WD_CELL_VERTICAL_ALIGNMENT.CENTER,
    ):
        """Заполняет ячейку, задавая стиль и выравнивание."""
        cell.text = text
        cell.vertical_alignment = valign
        paragraph = cell.paragraphs[0]
        if paragraph.runs:
            run = paragraph.runs[0]
            run.bold = bold
            run.font.name = "Times New Roman"
            run.font.size = Pt(11)
        paragraph.alignment = align

    def add_row():
        """Добавляет строку и возвращает список ячеек."""
        return table.add_row().cells

    def merge_range(cells, start, end):
        """Объединяет диапазон ячеек в строке."""
        if end <= start:
            return cells[start]
        return cells[start].merge(cells[end])

    row = add_row()
    right_cell = merge_range(row, 8, 11)
    set_cell(
        right_cell,
        f"Приложение № {spec_context['appendix_no']}".strip(),
        align=WD_ALIGN_PARAGRAPH.RIGHT,
    )

    row = add_row()
    right_cell = merge_range(row, 8, 11)
    set_cell(
        right_cell,
        spec_context["supply_contract_line"],
        align=WD_ALIGN_PARAGRAPH.RIGHT,
    )

    row = add_row()
    center_cell = merge_range(row, 0, 11)
    set_cell(
        center_cell,
        spec_context["spec_line"],
        bold=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )

    row = add_row()
    center_cell = merge_range(row, 0, 11)
    set_cell(
        center_cell,
        spec_context["contract_line"],
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )

    header_row = add_row()
    set_cell(header_row[0], "№", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    name_cell = merge_range(header_row, 1, 2)
    set_cell(
        name_cell,
        "Наименование Товара",
        bold=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )
    gost_cell = merge_range(header_row, 3, 4)
    set_cell(gost_cell, "ГОСТ / ТУ", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell(
        header_row[5],
        "Ед. изм.",
        bold=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )
    set_cell(
        header_row[6],
        "Кол-во",
        bold=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )
    set_cell(
        header_row[7],
        "Цена за ед. товара без НДС, руб",
        bold=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )
    set_cell(
        header_row[8],
        "НДС, руб",
        bold=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )
    set_cell(
        header_row[9],
        "Цена за ед. товара с НДС, руб",
        bold=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )
    set_cell(
        header_row[10],
        "Сумма без НДС, руб",
        bold=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )
    set_cell(
        header_row[11],
        "Общая стоимость товара c НДС, руб",
        bold=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )

    number_row = add_row()
    numbering = ["1", "2", "2", "3", "3", "4", "5", "6", "7", "8", "9", "10"]
    for idx, num in enumerate(numbering):
        set_cell(
            number_row[idx],
            num,
            bold=True,
            align=WD_ALIGN_PARAGRAPH.CENTER,
        )

    for idx, row_data in enumerate(rows, start=1):
        row_cells = add_row()
        set_cell(row_cells[0], str(idx), align=WD_ALIGN_PARAGRAPH.CENTER)
        name_cell = merge_range(row_cells, 1, 2)
        set_cell(name_cell, row_data["name"], align=WD_ALIGN_PARAGRAPH.LEFT)
        gost_cell = merge_range(row_cells, 3, 4)
        set_cell(
            gost_cell,
            row_data.get("gost", ""),
            align=WD_ALIGN_PARAGRAPH.LEFT,
        )
        set_cell(
            row_cells[5],
            row_data.get("unit", ""),
            align=WD_ALIGN_PARAGRAPH.CENTER,
        )
        set_cell(
            row_cells[6],
            format_quantity(row_data.get("qty")),
            align=WD_ALIGN_PARAGRAPH.CENTER,
        )
        set_cell(
            row_cells[7],
            format_decimal(row_data.get("net")),
            align=WD_ALIGN_PARAGRAPH.RIGHT,
        )
        set_cell(
            row_cells[8],
            format_decimal(row_data.get("vat")),
            align=WD_ALIGN_PARAGRAPH.RIGHT,
        )
        set_cell(
            row_cells[9],
            format_decimal(row_data.get("gross")),
            align=WD_ALIGN_PARAGRAPH.RIGHT,
        )
        set_cell(
            row_cells[10],
            format_decimal(row_data.get("sum_net")),
            align=WD_ALIGN_PARAGRAPH.RIGHT,
        )
        set_cell(
            row_cells[11],
            format_decimal(row_data.get("sum_gross")),
            align=WD_ALIGN_PARAGRAPH.RIGHT,
        )

    total_net = _sum_rows(rows, "sum_net")
    total_gross = _sum_rows(rows, "sum_gross")
    total_vat = quantize_money(total_gross - total_net)

    total_row = add_row()
    total_cell = merge_range(total_row, 0, 9)
    set_cell(total_cell, "Всего", bold=True, align=WD_ALIGN_PARAGRAPH.LEFT)
    set_cell(
        total_row[10],
        format_decimal(total_net),
        bold=True,
        align=WD_ALIGN_PARAGRAPH.RIGHT,
    )
    set_cell(
        total_row[11],
        format_decimal(total_gross),
        bold=True,
        align=WD_ALIGN_PARAGRAPH.RIGHT,
    )

    vat_row = add_row()
    vat_cell = merge_range(vat_row, 0, 10)
    set_cell(
        vat_cell,
        "В том числе НДС",
        bold=True,
        align=WD_ALIGN_PARAGRAPH.LEFT,
    )
    set_cell(
        vat_row[11],
        format_decimal(total_vat),
        bold=True,
        align=WD_ALIGN_PARAGRAPH.RIGHT,
    )

    for label, value in spec_context["detail_rows"]:
        row = add_row()
        label_cell = merge_range(row, 0, 1)
        value_cell = merge_range(row, 2, 11)
        set_cell(label_cell, label, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT)
        set_cell(value_cell, value, align=WD_ALIGN_PARAGRAPH.LEFT)

    row = add_row()
    buyer_cell = merge_range(row, 0, 5)
    seller_cell = merge_range(row, 6, 11)
    set_cell(
        buyer_cell,
        "Покупатель:",
        bold=True,
        align=WD_ALIGN_PARAGRAPH.LEFT,
    )
    set_cell(
        seller_cell,
        "Поставщик:",
        bold=True,
        align=WD_ALIGN_PARAGRAPH.LEFT,
    )

    row = add_row()
    buyer_cell = merge_range(row, 0, 5)
    seller_cell = merge_range(row, 6, 11)
    set_cell(
        buyer_cell,
        spec_context["buyer_position"],
        align=WD_ALIGN_PARAGRAPH.LEFT,
    )
    set_cell(
        seller_cell,
        spec_context["seller_position"],
        align=WD_ALIGN_PARAGRAPH.LEFT,
    )

    row = add_row()
    buyer_cell = merge_range(row, 0, 5)
    seller_cell = merge_range(row, 6, 11)
    set_cell(
        buyer_cell,
        spec_context["buyer_org"],
        align=WD_ALIGN_PARAGRAPH.LEFT,
    )
    set_cell(
        seller_cell,
        spec_context["seller_org"],
        align=WD_ALIGN_PARAGRAPH.LEFT,
    )

    row = add_row()
    buyer_cell = merge_range(row, 0, 5)
    seller_cell = merge_range(row, 6, 11)
    set_cell(
        buyer_cell,
        spec_context["buyer_sign"],
        align=WD_ALIGN_PARAGRAPH.LEFT,
    )
    set_cell(
        seller_cell,
        spec_context["seller_sign"],
        align=WD_ALIGN_PARAGRAPH.LEFT,
    )

    doc.save(path)


def _sum_rows(rows, key):
    """Суммирует значения Decimal по ключу."""
    total = Decimal("0")
    for row in rows:
        value = row.get(key)
        if value is not None:
            total += value
    return quantize_money(total)
