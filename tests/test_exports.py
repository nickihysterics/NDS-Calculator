from decimal import Decimal

from docx import Document
from openpyxl import Workbook, load_workbook

from calculations import normalize_rows
from excel_io import read_excel, write_excel
from settings_store import default_settings
from word_export import write_word


def sample():
    return normalize_rows(
        {
            "vatRate": "22",
            "rows": [
                {
                    "name": "Бумага",
                    "net": "350",
                    "qty": "5",
                    "unit": "упак",
                    "gost": "Учебный",
                    "source": "net",
                },
                {"name": '=HYPERLINK("https://example.invalid")', "gross": "610", "qty": "3"},
            ],
        }
    )


def test_excel_roundtrip_preserves_products_quantities_and_totals(tmp_path):
    rate, original = sample()
    path = tmp_path / "spec.xlsx"
    write_excel(path, rate, original)
    imported = read_excel(path)
    assert len(imported) == 2  # «Итого» не превращается в товар.
    _, restored = normalize_rows({"vatRate": str(rate), "rows": imported})
    for before, after in zip(original, restored):
        for key in ("name", "unit", "gost", "qty", "net", "gross", "sum_net", "sum_gross"):
            assert before[key] == after[key]
    workbook = load_workbook(path)
    sheet = workbook.active
    assert sheet["B3"].data_type == "s"
    assert sheet["J4"].value == 3965
    workbook.close()


def test_headerless_and_offset_header(tmp_path):
    for prefix in ([], [["Учебная спецификация"], ["Product name", "Net price", "Quantity"]]):
        workbook = Workbook()
        for row in prefix:
            workbook.active.append(row)
        workbook.active.append(["Бумага", 100, 2 if prefix else None])
        path = tmp_path / "input.xlsx"
        workbook.save(path)
        rows = read_excel(path)
        assert len(rows) == 1
        assert rows[0]["name"] == "Бумага"
        assert Decimal(rows[0]["net"]) == 100
        if prefix:
            assert rows[0]["qty"] == "2"


def test_word_totals_settings_and_landscape(tmp_path):
    rate, rows = sample()
    path = tmp_path / "spec.docx"
    settings = default_settings()
    settings.update(
        spec_date="14.09.2026", spec_line="Спецификация от __DATE__", buyer_org="Учебный покупатель"
    )
    write_word(path, rate, rows, settings)
    document = Document(path)
    text = "\n".join(cell.text for row in document.tables[0].rows for cell in row.cells)
    assert "3 965,00" in text
    assert "715,00" in text
    assert "Учебный покупатель" in text
    assert "Спецификация от 14.09.2026" in text
    assert "Бумага" in text
    assert document.sections[0].page_width > document.sections[0].page_height
